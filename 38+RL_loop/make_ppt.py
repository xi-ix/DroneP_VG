#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "38+RL_loop/38+RL_loop模型结构图.pptx"

COLORS = {
    "bg": RGBColor(247, 248, 250),
    "ink": RGBColor(31, 41, 55),
    "muted": RGBColor(82, 96, 113),
    "blue": RGBColor(220, 235, 255),
    "blue_line": RGBColor(79, 126, 201),
    "green": RGBColor(223, 244, 231),
    "green_line": RGBColor(76, 155, 104),
    "yellow": RGBColor(255, 240, 199),
    "yellow_line": RGBColor(196, 145, 47),
    "red": RGBColor(255, 225, 225),
    "red_line": RGBColor(198, 94, 94),
    "purple": RGBColor(238, 230, 255),
    "purple_line": RGBColor(124, 97, 201),
    "gray": RGBColor(237, 240, 243),
    "gray_line": RGBColor(154, 164, 178),
    "white": RGBColor(255, 255, 255),
    "dark": RGBColor(36, 52, 71),
}


def add_title(slide, title: str):
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.48))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = COLORS["dark"]


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = COLORS["bg"]


def textbox(slide, x, y, w, h, text, size=16, bold=False, color="ink", align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = COLORS[color]
    return shape


def box(slide, x, y, w, h, text, fill="blue", size=14, bold=False, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[fill]
    shape.line.color.rgb = COLORS.get(f"{fill}_line", COLORS["gray_line"])
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = COLORS["ink"]
    return shape


def arrow(slide, x1, y1, x2, y2, color="gray_line"):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = COLORS[color]
    conn.line.width = Pt(2)
    conn.line.end_arrowhead = True
    return conn


def metric_bar(slide):
    box(slide, 0.55, 6.72, 12.25, 0.48, "测试结果：38+RL_loop  mAP@0.5=0.1384  Acc@0.5=0.5161  Acc@0.75=0.3261  Pred=38370  |  STOP=0，全部样本进入第二轮闭环", "gray", 12, True)


def new_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, title)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = new_slide(prs, "38+RL_loop：强化学习驱动的闭环推理结构")
    textbox(s, 0.75, 0.95, 11.8, 0.8, "Query-aware 候选扩充 + RL 动态动作选择 + Exp38 多元元信息重排", 25, True, "dark")
    box(s, 1.0, 2.0, 3.0, 1.1, "输入\n图像 + 语言 query", "blue", 18, True)
    box(s, 5.15, 2.0, 3.0, 1.1, "闭环控制器\nRL policy", "purple", 18, True)
    box(s, 9.25, 2.0, 3.0, 1.1, "最终输出\n检测框 + 分数", "green", 18, True)
    arrow(s, 4.05, 2.55, 5.05, 2.55)
    arrow(s, 8.2, 2.55, 9.15, 2.55)
    box(s, 1.0, 4.0, 11.25, 1.0, "推理阶段真实闭环：第一轮候选 -> 状态反馈 -> RL 选择第二轮 prompt -> 候选合并 -> Exp38 重排", "yellow", 18, True)
    metric_bar(s)

    s = new_slide(prs, "总模块示意图")
    xs = [0.35, 2.4, 4.45, 6.5, 8.55, 10.6]
    texts = ["Image\n+ Query", "Query-aware\n第一轮候选 C0", "State\n38维反馈", "RL Policy\n12动作", "第二轮\nPrompt候选", "Exp38\n重排输出"]
    fills = ["blue", "blue", "yellow", "purple", "purple", "green"]
    for x, t, f in zip(xs, texts, fills):
        box(s, x, 1.1, 1.65, 1.0, t, f, 13, True)
    for i in range(len(xs) - 1):
        arrow(s, xs[i] + 1.68, 1.6, xs[i + 1] - 0.05, 1.6)
    box(s, 2.15, 3.15, 2.25, 0.85, "基础候选\nquery_alias_fallback", "blue", 14, True)
    box(s, 4.8, 3.05, 3.0, 1.05, "反馈特征\n数量 / 置信度 / 面积 / 类别 / 小目标比例", "yellow", 13)
    box(s, 8.25, 3.0, 2.5, 1.1, "动作执行\n从 Exp38 prompt 来源取候选", "purple", 13)
    arrow(s, 3.25, 2.1, 3.25, 3.1)
    arrow(s, 6.25, 2.1, 6.25, 3.0)
    arrow(s, 9.5, 2.1, 9.5, 2.95)
    box(s, 1.0, 5.0, 11.25, 0.8, "闭环不是训练时概念：test.py 推理时逐图执行 state -> action -> candidate expansion。", "red", 17, True)
    metric_bar(s)

    s = new_slide(prs, "模块1：输入 Query 与第一轮候选生成")
    box(s, 0.55, 1.05, 2.2, 0.9, "输入\nimage + language query", "blue", 15, True)
    arrow(s, 2.78, 1.5, 3.45, 1.5)
    box(s, 3.5, 0.85, 2.45, 1.3, "手写词表解析\n目标类 + alias\n不额外引入语言模型", "yellow", 14)
    arrow(s, 5.98, 1.5, 6.65, 1.5)
    box(s, 6.7, 0.85, 2.75, 1.3, "query_alias_fallback\n固定10类 + query类 + alias兜底", "blue", 14, True)
    arrow(s, 9.48, 1.5, 10.15, 1.5)
    box(s, 10.2, 1.05, 2.2, 0.9, "第一轮候选\nC0", "green", 16, True)
    box(s, 0.85, 3.1, 3.25, 1.15, "固定10类\npedestrian / people / bicycle / car / van / truck / tricycle / awning tricycle / bus / motor", "gray", 11)
    box(s, 4.65, 3.1, 3.25, 1.15, "Query目标类\n例如 car / person / motor / bicycle\n由词表匹配得到", "gray", 12)
    box(s, 8.45, 3.1, 3.25, 1.15, "Alias兜底\nmotorcycle / motorbike / scooter\ncovered tricycle 等", "gray", 12)
    box(s, 1.55, 5.1, 10.15, 0.8, "这一版第一步已经加入 query，不再是纯 fixed10，因此后续 RL 是在 query-aware 初始候选上做闭环反馈。", "yellow", 16, True)
    metric_bar(s)

    s = new_slide(prs, "模块2：闭环状态 State 构造")
    box(s, 0.55, 1.05, 2.25, 0.9, "第一轮候选 C0", "blue", 16, True)
    arrow(s, 2.85, 1.5, 3.55, 1.5)
    box(s, 3.6, 0.9, 2.4, 1.2, "统计反馈\n不看GT，不看测试答案", "yellow", 15, True)
    arrow(s, 6.05, 1.5, 6.75, 1.5)
    box(s, 6.8, 0.9, 2.35, 1.2, "38维 state", "purple", 18, True)
    arrow(s, 9.2, 1.5, 9.9, 1.5)
    box(s, 9.95, 0.9, 2.65, 1.2, "输入 RL policy\n决定第二轮动作", "green", 15, True)
    labels = [
        ("全局质量", "候选总数、均值分数、最高分、分数方差"),
        ("几何覆盖", "平均面积、小目标候选比例"),
        ("类别分布", "10类候选数量：10维"),
        ("类别置信度", "10类最高分：10维"),
        ("类别尺度", "10类平均面积：10维"),
        ("预算标记", "round id + remaining budget"),
    ]
    xys = [(0.75, 3.0), (4.55, 3.0), (8.35, 3.0), (0.75, 4.6), (4.55, 4.6), (8.35, 4.6)]
    for (title, body), (x, y) in zip(labels, xys):
        box(s, x, y, 3.05, 1.05, f"{title}\n{body}", "gray", 12, True)
    metric_bar(s)

    s = new_slide(prs, "模块3：RL Policy 网络结构")
    nodes = [(0.55, "state\n38"), (2.8, "Linear\n38 -> 128"), (4.9, "LayerNorm\n+ ReLU"), (6.8, "Linear\n128 -> 64"), (8.9, "LayerNorm\n+ ReLU"), (10.8, "Actor logits\n64 -> 12")]
    for x, text in nodes:
        box(s, x, 1.0, 1.55 if x not in (4.9, 8.9, 10.8) else (1.35 if x != 10.8 else 1.75), 0.95, text, "purple" if x > 2 else "yellow", 13, True)
    for x1, x2 in [(2.15, 2.75), (4.38, 4.85), (6.28, 6.75), (8.38, 8.85), (10.28, 10.75)]:
        arrow(s, x1, 1.48, x2, 1.48)
    box(s, 2.8, 3.05, 1.55, 0.85, "Linear\n38 -> 64", "blue", 13, True)
    arrow(s, 4.38, 3.48, 4.95, 3.48)
    box(s, 5.0, 3.05, 1.2, 0.85, "ReLU", "blue", 13, True)
    arrow(s, 6.25, 3.48, 6.82, 3.48)
    box(s, 6.85, 3.05, 1.6, 0.85, "Value\n64 -> 1", "blue", 13, True)
    box(s, 1.25, 4.75, 4.9, 0.9, "训练损失\npolicy_loss + 0.5 * value_loss + 0.01 * entropy_loss", "gray", 14)
    box(s, 7.0, 4.75, 4.9, 0.9, "训练设备\ncuda:0，torch.cuda=True，300 epochs，best epoch=295", "gray", 14)
    metric_bar(s)

    s = new_slide(prs, "模块4：动作空间与第二轮 Prompt 扩充")
    box(s, 0.55, 1.0, 2.3, 0.95, "RL logits\nargmax", "purple", 16, True)
    arrow(s, 2.9, 1.48, 3.65, 1.48)
    box(s, 3.7, 0.9, 2.6, 1.15, "12个动作\nSTOP + 11个Exp38 prompt", "yellow", 15, True)
    arrow(s, 6.35, 1.48, 7.1, 1.48)
    box(s, 7.15, 0.9, 2.6, 1.15, "按 source_prompt\n从 metadata 取候选", "blue", 15, True)
    arrow(s, 9.8, 1.48, 10.55, 1.48)
    box(s, 10.6, 1.0, 2.05, 0.95, "第二轮候选\nC_action", "green", 15, True)
    actions = "ADD_GDINO_BASE | ADD_PERSON | ADD_PEOPLE | ADD_GROUP_OF_PEOPLE\nADD_TRICYCLE | ADD_COVERED_TRICYCLE | ADD_AWNING_TRICYCLE\nADD_MOTORCYCLE | ADD_MOTORBIKE | ADD_SCOOTER | ADD_BICYCLE"
    box(s, 0.9, 3.1, 11.4, 1.0, actions, "gray", 14)
    box(s, 1.0, 4.8, 3.2, 0.95, "本次测试动作\nADD_PERSON:108", "green", 15, True)
    box(s, 5.0, 4.8, 3.2, 0.95, "ADD_MOTORCYCLE:29\nADD_PEOPLE:8", "green", 15, True)
    box(s, 9.0, 4.8, 3.2, 0.95, "ADD_COVERED_TRICYCLE:5\nSTOP:0", "red", 15, True)
    metric_bar(s)

    s = new_slide(prs, "模块5：候选合并、去重与 Exp38 重排")
    box(s, 0.55, 1.0, 2.15, 0.9, "C0\nquery-aware基础候选", "blue", 14, True)
    box(s, 0.55, 2.4, 2.15, 0.9, "C_action\n第二轮补充候选", "purple", 14, True)
    arrow(s, 2.75, 1.45, 3.55, 2.0)
    arrow(s, 2.75, 2.85, 3.55, 2.2)
    box(s, 3.6, 1.55, 2.25, 1.05, "合并\nC = C0 ∪ C_action", "yellow", 15, True)
    arrow(s, 5.9, 2.08, 6.55, 2.08)
    box(s, 6.6, 1.55, 2.1, 1.05, "同类NMS去重\nIoU阈值 0.92", "yellow", 14, True)
    arrow(s, 8.75, 2.08, 9.35, 2.08)
    box(s, 9.4, 1.35, 2.65, 1.45, "Exp38 score映射\n同类 IoU>=0.90\n使用 final_score", "green", 14, True)
    box(s, 1.1, 4.15, 3.25, 1.05, "匹配成功\n用 rich metadata fusion 分数排序", "green", 14)
    box(s, 5.0, 4.15, 3.25, 1.05, "匹配失败\n保留原始候选分数", "gray", 14)
    box(s, 8.9, 4.15, 3.25, 1.05, "最终预测\n按分数排序输出txt", "blue", 14)
    metric_bar(s)

    s = new_slide(prs, "模块6：训练阶段详细流程")
    steps = [("1 读取val样本", "150张图 + GT + 候选文件"), ("2 生成基础候选", "query_alias_fallback -> C0"), ("3 枚举动作", "12个动作逐一生成 C0∪C_action"), ("4 计算奖励", "召回/小目标/精度提升 - 成本"), ("5 构建表", "每图得到 12维 reward table"), ("6 GPU训练", "Actor-Critic policy")]
    xs = [0.45, 2.5, 4.55, 6.6, 8.65, 10.7]
    for x, (a, b) in zip(xs, steps):
        box(s, x, 1.1, 1.65, 1.2, f"{a}\n{b}", "yellow" if "奖励" in a or "表" in a else "blue", 10.5, True)
    for i in range(len(xs) - 1):
        arrow(s, xs[i] + 1.68, 1.7, xs[i + 1] - 0.05, 1.7)
    box(s, 0.85, 3.15, 3.4, 1.0, "Reward重点\n降低STOP收益\n降低候选数量惩罚", "red", 14, True)
    box(s, 4.95, 3.15, 3.4, 1.0, "小目标强化\nsmall_recall权重提高\n类别 {1,2,3,7,8,10}", "green", 14, True)
    box(s, 9.05, 3.15, 3.4, 1.0, "训练产物\nrl_policy_exp38_reward.pt\ntrain_summary.json", "purple", 14, True)
    box(s, 1.55, 5.15, 10.05, 0.75, "训练确认：torch=2.6.0+cu124，cuda=True，device=cuda:0，best epoch=295，val_mean_reward=0.064507", "gray", 14, True)
    metric_bar(s)

    s = new_slide(prs, "模块7：运行时推理与最终生成")
    steps = [("输入", "image + query"), ("Round 1", "query_alias_fallback\n生成C0"), ("反馈", "提取38维state"), ("策略", "RL policy选择动作"), ("Round 2", "追加Exp38 prompt候选"), ("输出", "合并去重\nExp38重排")]
    for idx, (a, b) in enumerate(steps):
        x = 0.55 + idx * 2.05
        box(s, x, 1.1, 1.65, 1.25, f"{a}\n{b}", ["blue", "blue", "yellow", "purple", "purple", "green"][idx], 10.5, True)
        if idx < len(steps) - 1:
            arrow(s, x + 1.68, 1.72, x + 2.0, 1.72)
    box(s, 0.95, 3.3, 3.0, 1.0, "闭环证据\n测试阶段逐图写入 closed_loop_trace.json", "yellow", 14, True)
    box(s, 5.0, 3.3, 3.0, 1.0, "动作分布\nSTOP=0，每张图均进入第二轮", "red", 14, True)
    box(s, 9.05, 3.3, 3.0, 1.0, "最终文件\npredictions/*.txt\nsummary.md/json", "green", 14, True)
    box(s, 1.35, 5.2, 10.6, 0.85, "最终性能：比纯RL_loop提升 mAP@0.5 0.1265 -> 0.1384；比query-base提升 Acc@0.5 0.4651 -> 0.5161", "gray", 15, True)
    metric_bar(s)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
