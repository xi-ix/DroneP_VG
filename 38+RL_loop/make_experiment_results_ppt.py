#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "PPT/三次实验结果汇总.pptx"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
C = {
    "bg": RGBColor(246, 247, 249),
    "paper": RGBColor(255, 255, 255),
    "ink": RGBColor(25, 36, 48),
    "muted": RGBColor(91, 104, 117),
    "line": RGBColor(218, 223, 229),
    "navy": RGBColor(30, 61, 89),
    "blue": RGBColor(45, 112, 170),
    "blue_light": RGBColor(225, 239, 250),
    "teal": RGBColor(23, 133, 125),
    "teal_light": RGBColor(222, 243, 239),
    "amber": RGBColor(217, 145, 36),
    "amber_light": RGBColor(251, 238, 211),
    "coral": RGBColor(198, 80, 73),
    "gray": RGBColor(181, 190, 199),
    "gray_light": RGBColor(236, 239, 242),
}


def rect(slide, x, y, w, h, fill="paper", line="line", radius=0.06):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = C[fill]
    s.line.color.rgb = C[line]
    s.line.width = Pt(0.8)
    if radius and hasattr(s.adjustments, "__len__") and len(s.adjustments):
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, value, size=12, color="ink", bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.03):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = value
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = C[color]
    return s


def rich_text(slide, x, y, w, h, runs, size=12, align=PP_ALIGN.LEFT):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    for value, color, bold in runs:
        r = p.add_run()
        r.text = value
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = C[color]
    return s


def base_slide(prs, number, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["bg"]
    rect(s, 0, 0, W, 0.11, "navy", "navy", 0)
    text(s, 0.52, 0.28, 0.55, 0.34, f"0{number}", 12, "blue", True)
    text(s, 1.02, 0.22, 11.75, 0.52, title, 24, "ink", True)
    text(s, 1.03, 0.72, 11.7, 0.30, subtitle, 10.5, "muted")
    text(s, 12.46, 7.12, 0.38, 0.20, str(number), 9, "muted", align=PP_ALIGN.RIGHT)
    return s


def section_label(slide, x, y, w, value, color="blue"):
    rect(slide, x, y + 0.07, 0.06, 0.23, color, color, 0)
    text(slide, x + 0.14, y, w - 0.14, 0.36, value, 11, "ink", True)


def metric_chip(slide, x, y, w, label, value, accent="blue", value_size=17):
    rect(slide, x, y, w, 0.72, "paper", "line")
    rect(slide, x, y, 0.06, 0.72, accent, accent, 0)
    text(slide, x + 0.18, y + 0.08, w - 0.27, 0.22, label, 9.5, "muted")
    text(slide, x + 0.18, y + 0.30, w - 0.27, 0.31, value, value_size, accent, True)


def table(slide, x, y, widths, row_h, headers, rows, highlights=None, font=10.2):
    highlights = highlights or {}
    total_w = sum(widths)
    rect(slide, x, y, total_w, row_h * (len(rows) + 1), "paper", "line")
    cx = x
    for i, (head, width) in enumerate(zip(headers, widths)):
        rect(slide, cx, y, width, row_h, "navy", "navy", 0)
        text(slide, cx + 0.02, y + 0.02, width - 0.04, row_h - 0.04, head,
             9.2, "paper", True, PP_ALIGN.CENTER)
        cx += width
    for r_idx, row in enumerate(rows):
        cy = y + row_h * (r_idx + 1)
        fill = highlights.get(r_idx, "paper")
        cx = x
        for c_idx, (value, width) in enumerate(zip(row, widths)):
            rect(slide, cx, cy, width, row_h, fill, "line", 0)
            align = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
            color = "blue" if r_idx in highlights and c_idx in (0, len(row) - 1) else "ink"
            text(slide, cx + 0.04, cy + 0.01, width - 0.08, row_h - 0.02, str(value),
                 font, color, r_idx in highlights, align)
            cx += width


def draw_line_chart(slide, x, y, w, h, rounds, acc, latency):
    # Two normalized series share the plot; exact values are labelled at every point.
    rect(slide, x, y, w, h, "paper", "line")
    text(slide, x + 0.25, y + 0.15, w - 0.5, 0.30, "轮数增加带来的收益与时延", 12, "ink", True)
    px, py, pw, ph = x + 0.62, y + 0.72, w - 0.95, h - 1.15
    for i in range(4):
        yy = py + ph * i / 3
        rect(slide, px, yy, pw, 0.006, "line", "line", 0)
    for i, r in enumerate(rounds):
        xx = px + pw * i / (len(rounds) - 1)
        text(slide, xx - 0.22, py + ph + 0.10, 0.44, 0.24, f"{r}轮", 9.5, "muted", align=PP_ALIGN.CENTER)
    series = [(acc, min(acc) - 0.006, max(acc) + 0.006, "blue"),
              (latency, min(latency) - 1.0, max(latency) + 1.0, "amber")]
    for values, lo, hi, color in series:
        points = []
        for i, value in enumerate(values):
            xx = px + pw * i / (len(values) - 1)
            yy = py + ph * (1 - (value - lo) / (hi - lo))
            points.append((xx, yy))
        for (x1, y1), (x2, y2) in zip(points, points[1:]):
            line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            line.line.color.rgb = C[color]
            line.line.width = Pt(2.5)
        for idx, ((xx, yy), value) in enumerate(zip(points, values)):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xx - 0.065), Inches(yy - 0.065), Inches(0.13), Inches(0.13))
            dot.fill.solid(); dot.fill.fore_color.rgb = C[color]; dot.line.color.rgb = C[color]
            label = f"{value:.4f}" if color == "blue" else f"{value:.2f}s"
            yoff = -0.34 if color == "blue" else 0.12
            text(slide, xx - 0.36, yy + yoff, 0.72, 0.24, label, 9.5, color, True, PP_ALIGN.CENTER)
    rich_text(slide, x + w - 2.55, y + 0.16, 2.20, 0.24,
              [("● ", "blue", True), ("Acc@0.5", "muted", False), ("    ● ", "amber", True), ("闭环时间", "muted", False)], 9.5, PP_ALIGN.CENTER)


def draw_horizontal_bars(slide, x, y, w, h, labels, values, colors, title, value_fmt="{:.4f}"):
    rect(slide, x, y, w, h, "paper", "line")
    text(slide, x + 0.25, y + 0.14, w - 0.5, 0.32, title, 12, "ink", True)
    max_v = max(values) * 1.12
    row_h = (h - 0.72) / len(labels)
    label_w = 1.35
    for i, (label, value, color) in enumerate(zip(labels, values, colors)):
        cy = y + 0.55 + i * row_h
        text(slide, x + 0.14, cy, label_w, row_h * 0.85, label, 9.3, "muted", align=PP_ALIGN.RIGHT)
        bw = (w - label_w - 1.10) * value / max_v
        rect(slide, x + label_w + 0.24, cy + row_h * 0.24, bw, row_h * 0.40, color, color, 0)
        text(slide, x + label_w + 0.31 + bw, cy, 0.72, row_h * 0.85, value_fmt.format(value), 9.2, color, True)


def slide_rounds(prs):
    s = base_slide(prs, 1, "RL 闭环轮数消融", "固定同一策略与候选流程 · VisDrone test：150 张 / 7,973 GT · 仅改变总推理轮数")
    metric_chip(s, 0.55, 1.18, 2.10, "最高 Acc@0.5", "0.5357 · 4轮", "blue")
    metric_chip(s, 2.78, 1.18, 2.10, "最高 mAP@0.5", "0.1402 · 4轮", "teal")
    metric_chip(s, 5.01, 1.18, 2.10, "均衡设置", "3轮 · 93ms/图", "amber", 14.5)
    draw_line_chart(s, 7.35, 1.18, 5.43, 3.31, [2, 3, 4], [0.5139, 0.5293, 0.5357], [11.27, 13.97, 15.92])
    section_label(s, 0.55, 2.13, 6.55, "同口径复跑结果")
    table(s, 0.55, 2.55, [0.75, 1.10, 1.05, 1.05, 1.05, 1.15], 0.46,
          ["轮数", "候选框", "TP@0.5", "Acc@0.5", "mAP@0.5", "单图时间"],
          [["2", "37,487", "4,097", "0.5139", "0.1386", "75.11 ms"],
           ["3", "40,904", "4,220", "0.5293", "0.1397", "93.16 ms"],
           ["4", "42,229", "4,271", "0.5357", "0.1402", "106.11 ms"]],
          {1: "amber_light", 2: "blue_light"}, 9.6)
    rect(s, 0.55, 4.62, 12.23, 1.70, "paper", "line")
    section_label(s, 0.82, 4.82, 3.2, "结论")
    rich_text(s, 0.88, 5.18, 11.60, 0.40,
              [("4轮", "blue", True), ("取得最高指标，但 3→4 轮仅新增 ", "ink", False),
               ("51 TP", "blue", True), ("、mAP +0.00052，同时增加 ", "ink", False),
               ("1.94 s", "amber", True), (" 推理时间。", "ink", False)], 15)
    text(s, 0.88, 5.66, 11.55, 0.36,
         "最高召回/评分优先：4轮；综合候选规模与推理效率：3轮更均衡。", 13, "muted")
    text(s, 0.58, 6.76, 11.8, 0.22, "注：4轮沿用三轮训练策略，属于固定策略下的推理轮数消融。", 9.2, "muted")


def slide_actions(prs):
    s = base_slide(prs, 2, "RL 动作策略消融（Exp43）", "相同初始候选、去重、Exp38 重评分与评测逻辑 · Random action 为 5 seeds 均值")
    metric_chip(s, 0.55, 1.18, 2.30, "Fixed RL vs Random", "Acc +4.01 pp", "blue")
    metric_chip(s, 2.98, 1.18, 2.30, "小目标最终召回", "+5.75 pp", "teal")
    metric_chip(s, 5.41, 1.18, 2.30, "相对穷举动作预算", "仅 1 / 8.13", "amber")
    draw_horizontal_bars(s, 7.94, 1.18, 4.84, 3.73,
                         ["No expansion", "Random", "Prompt-rank", "Fixed RL", "Qwen dynamic", "All expansion"],
                         [0.4651, 0.4892, 0.5090, 0.5293, 0.5004, 0.5731],
                         ["gray", "gray", "teal", "blue", "amber", "coral"],
                         "Acc@0.5 对比")
    section_label(s, 0.55, 2.13, 7.15, "准确率、排序质量与资源开销")
    table(s, 0.55, 2.55, [1.55, 1.12, 1.02, 1.02, 1.06, 1.18], 0.39,
          ["方法", "预测框", "Acc@0.5", "mAP@0.5", "动作/图", "小目标Recall"],
          [["No expansion", "22,298", "0.4651", "0.1343", "0.000", "0.3291"],
           ["All expansion", "68,771", "0.5731", "0.1510", "11.000", "0.4593"],
           ["Random (5 seeds)", "30,227", "0.4892", "0.1377", "1.353", "0.3584"],
           ["Prompt-rank", "41,969", "0.5090", "0.1440", "1.087", "0.3840"],
           ["Fixed-action RL", "40,904", "0.5293", "0.1397", "1.353", "0.4159"],
           ["Qwen dynamic", "35,869", "0.5004", "0.1424", "0.953", "0.3732"]],
          {4: "blue_light"}, 8.7)
    rect(s, 0.55, 5.44, 12.23, 1.05, "paper", "line")
    section_label(s, 0.82, 5.59, 2.0, "结论")
    rich_text(s, 2.05, 5.55, 10.35, 0.42,
              [("Fixed-action RL", "blue", True), (" 在相同动作预算下显著优于随机策略，证明收益来自", "ink", False),
               ("有效的动作选择", "teal", True), ("。", "ink", False)], 13.2)
    text(s, 2.05, 5.98, 10.35, 0.30,
         "All expansion 绝对指标最高但开销最大；动态策略更省候选，固定策略在召回上更强。", 11.5, "muted")
    text(s, 0.58, 6.76, 11.8, 0.22, "Small-object：面积 ≤ 32×32，class-aware 一对一匹配，IoU ≥ 0.5。", 9.2, "muted")


def slide_metadata(prs):
    s = base_slide(prs, 3, "Exp38 元数据特征消融", "固定 52,872 个候选框与 19,649 参数 · 3 seeds · alpha 仅由 val holdout 选择")
    metric_chip(s, 0.55, 1.18, 2.35, "Full vs Raw mAP", "+0.0212", "blue")
    metric_chip(s, 3.03, 1.18, 2.35, "相对提升", "+14.5%", "teal")
    metric_chip(s, 5.51, 1.18, 2.35, "最大单步贡献", "+source · +0.0150", "amber")
    draw_horizontal_bars(s, 8.09, 1.18, 4.69, 4.24,
                         ["Raw score", "+ geometry", "+ query", "+ source", "+ small prior", "Full Exp38"],
                         [0.1464, 0.1504, 0.1514, 0.1664, 0.1637, 0.1677],
                         ["gray", "gray", "teal", "amber", "coral", "blue"],
                         "mAP@0.5 逐步累积", "{:.4f}")
    section_label(s, 0.55, 2.13, 7.30, "逐步累积特征结果")
    table(s, 0.55, 2.55, [1.52, 0.80, 1.17, 1.17, 1.17, 1.22], 0.41,
          ["配置", "维数", "Acc@0.5", "Acc@0.75", "mAP@0.5", "逐步增量"],
          [["Raw score", "1", "0.5374", "0.3463", "0.1464", "+0.0000"],
           ["+ geometry", "9", "0.5374", "0.3463", "0.1504", "+0.0040"],
           ["+ query", "10", "0.5374", "0.3461", "0.1514", "+0.0010"],
           ["+ source", "36", "0.5374", "0.3483", "0.1664", "+0.0150"],
           ["+ small prior", "59", "0.5374", "0.3482", "0.1637", "−0.0027"],
           ["Full Exp38", "68", "0.5374", "0.3490", "0.1677", "+0.0040"]],
          {3: "amber_light", 5: "blue_light"}, 9.0)
    rect(s, 0.55, 5.53, 12.23, 0.98, "paper", "line")
    section_label(s, 0.82, 5.67, 2.0, "结论")
    rich_text(s, 2.05, 5.61, 10.35, 0.40,
              [("候选数和 Acc@0.5 完全不变", "ink", True), ("，mAP 的提升来自", "ink", False),
               ("排序质量", "blue", True), ("；", "ink", False), ("source 信息", "amber", True),
               ("是最主要且稳定的贡献。", "ink", False)], 13.0)
    text(s, 2.05, 6.02, 10.30, 0.25, "small prior 单独加入后下降，提示手工先验可能过拟合或与来源特征重复。", 11.2, "muted")
    text(s, 0.58, 6.76, 11.8, 0.22, "结果为均值；原始报告同时给出样本标准差。", 9.2, "muted")


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "三次实验结果汇总"
    prs.core_properties.subject = "RL 闭环轮数、动作策略与 Exp38 元数据消融"
    prs.core_properties.author = "DroneP_VG"
    slide_rounds(prs)
    slide_actions(prs)
    slide_metadata(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
