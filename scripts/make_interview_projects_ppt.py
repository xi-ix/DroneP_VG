#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "PPT"
OUT_PPTX = OUT_DIR / "面试项目经历概览.pptx"
OUT_SCRIPT = OUT_DIR / "面试项目经历概览_讲稿.md"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"

C = {
    "bg": RGBColor(247, 248, 250),
    "paper": RGBColor(255, 255, 255),
    "ink": RGBColor(26, 35, 46),
    "muted": RGBColor(88, 99, 112),
    "line": RGBColor(218, 223, 229),
    "navy": RGBColor(35, 61, 82),
    "blue": RGBColor(43, 112, 171),
    "blue_light": RGBColor(231, 241, 249),
    "teal": RGBColor(27, 130, 120),
    "teal_light": RGBColor(228, 244, 240),
    "amber": RGBColor(206, 132, 35),
    "amber_light": RGBColor(252, 241, 218),
    "coral": RGBColor(190, 80, 71),
    "coral_light": RGBColor(250, 234, 232),
}

PROJECTS = [
    {
        "number": "01",
        "tag": "视觉感知",
        "title": "无人机编队状态识别",
        "body": (
            "面向无人机集群飞行监测场景，使用 YOLO 对编队飞行视频进行逐帧目标检测，"
            "结合无人机之间的空间分布与形状特征，识别编队队形并判断其飞行状态，"
            "完成从视频输入、目标检测到状态输出的视觉分析流程。"
        ),
        "accent": "blue",
        "light": "blue_light",
    },
    {
        "number": "02",
        "tag": "语言智能体",
        "title": "交通数据 NLP2SQL Agent",
        "body": (
            "面向深圳道路交通检测平台，构建自然语言查询数据库的 Agent。"
            "系统能够理解用户对道路流量、检测记录等交通数据的查询意图，"
            "自动生成并执行 SQL，再将数据库结果整理为可读反馈，降低专业数据的使用门槛。"
        ),
        "accent": "teal",
        "light": "teal_light",
    },
    {
        "number": "03",
        "tag": "视觉语言",
        "title": "航拍开放词汇目标检测",
        "body": (
            "研究根据自然语言描述在无人机航拍图像中定位目标。针对小目标难识别、"
            "相似类别易混淆等问题，以 GroundingDINO 为基础，围绕提示词优化、"
            "候选框扩展与动态重排序进行改进，提升模型对开放类别目标的理解与定位能力。"
        ),
        "accent": "amber",
        "light": "amber_light",
    },
    {
        "number": "04",
        "tag": "具身智能",
        "title": "交通指挥机器人",
        "body": (
            "基于众擎机器人开发面向道路场景的交通指挥系统，使机器人能够自主转向、"
            "感知交通环境并通过动作进行指挥。本人负责视觉识别模块，重点检测未佩戴头盔、"
            "闯红灯等交通违法行为，为机器人后续决策与动作执行提供感知信息。"
        ),
        "accent": "coral",
        "light": "coral_light",
    },
]

SPEAKER_NOTES = """我的项目经历主要围绕人工智能在真实场景中的感知、理解和决策展开。

在视觉方向，我做过基于 YOLO 的无人机编队状态识别，通过飞行视频识别编队形状和状态；在语言模型方向，我开发过面向深圳道路交通检测平台的 NLP2SQL Agent，让用户可以直接用自然语言查询交通数据。

我的毕业设计研究无人机航拍场景下的开放词汇目标检测，重点解决小目标难定位和相似类别易混淆的问题。目前我参与具身智能交通指挥机器人项目，主要负责未佩戴头盔、闯红灯等交通违法行为的视觉识别。

这些项目覆盖了计算机视觉、视觉语言模型、智能体和具身智能，也让我逐步形成了从环境感知、信息理解到决策执行的完整技术视角。"""

FUTURE_NOTES = """未来我希望重点研究视觉感知与具身智能的结合。相比只在静态数据集上完成识别，我更关注智能体如何在真实、动态的环境中持续感知，并依据视觉信息作出可靠行动。

具体来说，我计划从三个方向展开。第一是可靠的动态视觉感知，继续研究目标检测、跟踪和行为识别，重点关注小目标、遮挡、视角变化以及复杂环境干扰，使机器人获得稳定、连续的环境状态。第二是视觉语言理解，希望利用视觉语言模型把图像、自然语言指令和场景知识联系起来，让机器人不仅能看见目标，还能理解人的任务要求和目标之间的关系。第三是感知与决策闭环，把视觉结果转化为可执行状态，并结合策略学习、强化学习等方法，让机器人根据环境反馈调整动作，而不是停留在单次识别上。

在研究生阶段，我希望先夯实计算机视觉、机器人学和强化学习基础，复现并理解具身智能领域的代表性工作；之后依托交通指挥机器人等真实平台，逐步完成从感知、理解、决策到实机验证的完整系统。我最终希望提升机器人在开放环境中的泛化能力、可靠性和实际应用价值。"""


def add_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C[fill]
    shape.line.color.rgb = C[line or fill]
    shape.line.width = Pt(0.8)
    if radius and len(shape.adjustments):
        shape.adjustments[0] = 0.06
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    value,
    size,
    color="ink",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(0)
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.12
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = C[color]
    return box


def add_project(slide, project, x, y, w, h):
    add_text(slide, x, y, 0.72, 0.54, project["number"], 25, project["accent"], True)
    add_text(slide, x + 0.77, y + 0.08, w - 0.77, 0.32, project["tag"], 10.2, "muted", True)
    add_text(slide, x, y + 0.82, w, 0.82, project["title"], 17, "ink", True)
    add_rect(slide, x, y + 1.77, w, 0.035, project["accent"], project["accent"], radius=False)
    add_text(slide, x, y + 2.10, w, h - 2.10, project["body"], 11.5, "muted")


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C["bg"]
    return slide


def build_title_slide(prs):
    slide = new_slide(prs)
    add_rect(slide, 0, 0, W, 0.07, "navy", "navy", radius=False)
    add_text(slide, 0.72, 0.78, 2.6, 0.30, "INTERVIEW PRESENTATION", 9.5, "muted", True)
    add_text(slide, 0.72, 2.08, 11.8, 0.86, "个人项目经历与未来规划", 31, "ink", True)
    add_text(slide, 0.74, 3.17, 9.7, 0.42, "从视觉感知、语言理解到具身智能的项目实践", 14, "muted")
    add_rect(slide, 0.74, 3.90, 1.55, 0.045, "blue", "blue", radius=False)
    add_rect(slide, 2.29, 3.90, 1.55, 0.045, "teal", "teal", radius=False)
    add_rect(slide, 3.84, 3.90, 1.55, 0.045, "amber", "amber", radius=False)
    add_rect(slide, 5.39, 3.90, 1.55, 0.045, "coral", "coral", radius=False)
    add_text(slide, 0.74, 6.63, 3.2, 0.30, "项目经历  ·  技术方向  ·  未来规划", 10.5, "navy", True)
    add_text(slide, 11.40, 6.63, 1.20, 0.30, "01 / 04", 9.5, "muted", True, PP_ALIGN.RIGHT)
    return slide


def build_projects_slide(prs):
    slide = new_slide(prs)

    add_rect(slide, 0, 0, W, 0.07, "navy", "navy", radius=False)
    add_text(slide, 0.58, 0.38, 8.4, 0.52, "项目经历与技术方向", 25, "ink", True)
    add_text(
        slide,
        0.57,
        0.94,
        11.9,
        0.34,
        "围绕真实场景，持续探索从视觉感知、语言理解到智能决策与动作执行的完整链路",
        11.5,
        "muted",
    )

    add_text(slide, 10.18, 0.45, 2.55, 0.30, "PROJECT PORTFOLIO", 9.5, "muted", True, PP_ALIGN.RIGHT)

    column_x = [0.58, 3.77, 6.96, 10.15]
    for index, (project, x) in enumerate(zip(PROJECTS, column_x)):
        add_project(slide, project, x, 1.58, 2.58, 4.62)
        if index < 3:
            add_rect(slide, x + 2.85, 1.62, 0.012, 4.72, "line", "line", radius=False)

    add_rect(slide, 0.58, 6.54, 12.17, 0.012, "line", "line", radius=False)
    add_text(
        slide,
        0.58,
        6.73,
        10.40,
        0.34,
        "技术关键词   计算机视觉  ·  目标检测  ·  视觉语言模型  ·  LLM Agent  ·  具身智能",
        11.2,
        "navy",
        True,
        PP_ALIGN.LEFT,
        MSO_ANCHOR.MIDDLE,
    )
    add_text(slide, 11.40, 6.73, 1.20, 0.30, "02 / 04", 9.5, "muted", True, PP_ALIGN.RIGHT)

    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.text = SPEAKER_NOTES
    return slide


def build_future_slide(prs):
    slide = new_slide(prs)
    add_rect(slide, 0, 0, W, 0.07, "navy", "navy", radius=False)
    add_text(slide, 0.62, 0.48, 1.00, 0.50, "03", 25, "teal", True)
    add_text(slide, 1.52, 0.51, 5.8, 0.48, "未来规划", 25, "ink", True)
    add_text(slide, 10.10, 0.58, 2.60, 0.28, "FUTURE PLAN", 9.5, "muted", True, PP_ALIGN.RIGHT)
    add_rect(slide, 0.62, 1.34, 12.08, 0.012, "line", "line", radius=False)

    add_text(
        slide,
        0.62,
        1.68,
        12.08,
        0.52,
        "研究主线：面向真实动态场景，构建“感知—理解—决策—行动”的具身智能闭环",
        16,
        "navy",
        True,
    )

    future_columns = [
        (
            "01",
            "可靠的动态视觉感知",
            "研究目标检测、目标跟踪与行为识别，重点应对小目标、遮挡、视角变化和复杂环境干扰，为机器人提供稳定、连续的环境状态。",
            "blue",
        ),
        (
            "02",
            "视觉语言与场景理解",
            "融合图像、自然语言指令和场景知识，使机器人从“识别目标”进一步走向理解任务意图、目标关系与环境语义。",
            "teal",
        ),
        (
            "03",
            "感知驱动的决策与行动",
            "将视觉结果转化为可执行状态，结合策略学习与强化学习，根据环境反馈持续调整动作，并在真实机器人平台完成验证。",
            "coral",
        ),
    ]
    column_x = [0.62, 4.58, 8.54]
    for index, ((number, title, body, accent), x) in enumerate(zip(future_columns, column_x)):
        add_text(slide, x, 2.56, 0.64, 0.46, number, 21, accent, True)
        add_text(slide, x + 0.69, 2.60, 3.05, 0.40, title, 14.5, "ink", True)
        add_rect(slide, x, 3.20, 3.55, 0.035, accent, accent, radius=False)
        add_text(slide, x, 3.52, 3.55, 1.72, body, 11.2, "muted")
        if index < 2:
            add_rect(slide, x + 3.74, 2.58, 0.012, 2.72, "line", "line", radius=False)

    add_rect(slide, 0.62, 5.64, 12.08, 0.012, "line", "line", radius=False)
    add_text(slide, 0.62, 5.91, 1.48, 0.34, "研究推进路径", 11, "navy", True)
    add_text(
        slide,
        2.15,
        5.89,
        10.00,
        0.38,
        "夯实视觉、机器人与强化学习基础   →   完成算法复现与系统构建   →   依托真实平台开展实机验证",
        10.8,
        "muted",
    )
    add_text(slide, 11.40, 6.63, 1.20, 0.30, "03 / 04", 9.5, "muted", True, PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = FUTURE_NOTES
    return slide


def build_end_slide(prs):
    slide = new_slide(prs)
    add_rect(slide, 0, 0, W, 0.07, "navy", "navy", radius=False)
    add_text(slide, 0.72, 2.23, 11.89, 0.86, "感谢各位老师聆听", 33, "ink", True, PP_ALIGN.CENTER)
    add_text(slide, 0.72, 3.31, 11.89, 0.42, "恳请各位老师批评指正", 14, "muted", False, PP_ALIGN.CENTER)
    add_rect(slide, 5.19, 4.18, 0.74, 0.04, "blue", "blue", radius=False)
    add_rect(slide, 5.93, 4.18, 0.74, 0.04, "teal", "teal", radius=False)
    add_rect(slide, 6.67, 4.18, 0.74, 0.04, "amber", "amber", radius=False)
    add_rect(slide, 7.41, 4.18, 0.74, 0.04, "coral", "coral", radius=False)
    add_text(slide, 0.72, 6.63, 2.0, 0.30, "THANK YOU", 9.5, "muted", True)
    add_text(slide, 11.40, 6.63, 1.20, 0.30, "04 / 04", 9.5, "muted", True, PP_ALIGN.RIGHT)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    build_title_slide(prs)
    build_projects_slide(prs)
    build_future_slide(prs)
    build_end_slide(prs)

    prs.core_properties.title = "个人项目经历与未来规划"
    prs.core_properties.subject = "面试展示"
    prs.core_properties.author = "Wangzhe"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    OUT_SCRIPT.write_text(
        "# 《个人项目经历与未来规划》讲稿\n\n"
        "## 第二页：项目经历与技术方向\n\n"
        + SPEAKER_NOTES
        + "\n\n## 第三页：未来规划\n\n"
        + FUTURE_NOTES
        + "\n",
        encoding="utf-8",
    )
    print(OUT_PPTX)
    print(OUT_SCRIPT)


if __name__ == "__main__":
    build()
