#!/usr/bin/env python3
from __future__ import annotations

import copy
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/wangzhe/.codex/attachments/7a6dfd99-87dc-476f-ae85-0e98fb2cfd9d/本科综合设计.docx")
OUT = ROOT / "论文/本科综合设计_最新方法与实验版.docx"
ASSET_DIR = ROOT / "report_assets/thesis_latest"
EXP38_FIG = ROOT / "report_assets/rl_loop_diagrams/exp38_rescoring_structure.png"
ACTION_FIG = ROOT / "experiment/exp43_rl_action_policy_ablation_20260911/log/figures/main_metrics_and_candidates.png"

BODY_STYLE = "论文正文-首行缩进"
SECTION_STYLE = "一级节标题2.3"
SUBSECTION_STYLE = "二级节标题2.3.1"
CAPTION_STYLE = "图题"
TABLE_CAPTION_STYLE = "表题注"


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"paragraph not found: {text}")


def find_heading(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1" and paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"chapter heading not found: {text}")


def remove_between(start, end):
    node = start._p.getnext()
    while node is not None and node is not end._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node


def add_p(anchor, value: str = "", style: str = BODY_STYLE, align=None, bold=False):
    p = anchor.insert_paragraph_before(style=style)
    if value:
        run = p.add_run(value)
        run.bold = bold
    if align is not None:
        p.alignment = align
    return p


def add_formula(anchor, value: str):
    p = add_p(anchor, value, BODY_STYLE, WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        run.font.name = "Cambria Math"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    return p


def set_cell_text(cell, value: str, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc: Document, anchor, caption: str, headers, rows, widths=None, highlight_rows=()):
    caption_p = add_p(anchor, caption, "Normal", WD_ALIGN_PARAGRAPH.CENTER)
    caption_p.paragraph_format.space_before = Pt(6)
    caption_p.paragraph_format.space_after = Pt(4)
    for run in caption_p.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(10.5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table._tbl.getparent().remove(table._tbl)
    anchor._p.addprevious(table._tbl)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=(255, 255, 255), size=8.3)
        shade(table.rows[0].cells[idx], "1E3D59")
    repeat_header(table.rows[0])
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, value in enumerate(values):
            set_cell_text(cells[c_idx], value, bold=r_idx in highlight_rows, size=8.1)
            if r_idx in highlight_rows:
                shade(cells[c_idx], "E1EFFA")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    add_p(anchor, "", BODY_STYLE)
    return table


def add_picture(doc: Document, anchor, path: Path, caption: str, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    p._p.getparent().remove(p._p)
    anchor._p.addprevious(p._p)
    caption_p = add_p(anchor, caption, "Normal", WD_ALIGN_PARAGRAPH.CENTER)
    caption_p.paragraph_format.space_before = Pt(4)
    caption_p.paragraph_format.space_after = Pt(6)
    for run in caption_p.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(10.5)


def draw_box(ax, xy, wh, title, detail, color):
    x, y = xy
    w, h = wh
    patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.012,rounding_size=0.018",
                           linewidth=1.5, edgecolor=color, facecolor="#FFFFFF")
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h * 0.64, title, ha="center", va="center", fontsize=13,
            fontweight="bold", color="#192430")
    ax.text(x + w / 2, y + h * 0.30, detail, ha="center", va="center", fontsize=9.5,
            color="#5B6875", linespacing=1.35)


def arrow(ax, p1, p2, color="#718096"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=14,
                                 linewidth=1.6, color=color))


def make_framework_figures():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    font_path = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    if Path(font_path).exists():
        font_manager.fontManager.addfont(font_path)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=font_path).get_name()
    plt.rcParams["axes.unicode_minus"] = False

    fig, ax = plt.subplots(figsize=(13.2, 5.2), dpi=180)
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.03, 0.92, "Qwen 动态提示词 + RL 闭环选择 + 富元信息融合排序总体流程",
            fontsize=19, fontweight="bold", color="#192430")
    xs = [0.03, 0.215, 0.40, 0.585, 0.77]
    titles = ["输入", "动态提示词池", "RL 闭环选择", "候选融合", "富元信息排序"]
    details = ["航拍图像 I\n文本查询 q", "Qwen2.5-VL-3B\n类别 + 短检测短语", "38维状态 + 20维动作\nSTOP / 扩展动作", "Grounding DINO\n合并 + 同类去重", "68维富元信息\n重标定 + NMS"]
    colors = ["#2D70AA", "#17857D", "#D99124", "#2D70AA", "#17857D"]
    for x, title, detail, color in zip(xs, titles, details, colors):
        draw_box(ax, (x, 0.47), (0.155, 0.27), title, detail, color)
    for x1, x2 in zip(xs, xs[1:]):
        arrow(ax, (x1 + 0.155, 0.605), (x2 - 0.01, 0.605))
    ax.add_patch(FancyArrowPatch((0.665, 0.42), (0.49, 0.42), connectionstyle="arc3,rad=-0.28",
                                 arrowstyle="-|>", mutation_scale=14, linewidth=1.5, color="#D99124"))
    ax.text(0.575, 0.26, "状态更新：候选数量、分数、面积、类别统计与剩余预算",
            ha="center", fontsize=10.2, color="#8B5A13")
    ax.text(0.50, 0.10, "训练阶段用检测增益构造奖励；推理阶段直接选择策略得分最高的有效动作",
            ha="center", fontsize=10.5, color="#5B6875")
    fig.tight_layout(pad=0.4)
    path = ASSET_DIR / "latest_overall_pipeline.png"
    fig.savefig(path, bbox_inches="tight", facecolor="#F6F7F9")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(12.8, 5.0), dpi=180)
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.03, 0.91, "动态策略网络：状态—动作匹配评分", fontsize=19, fontweight="bold", color="#192430")
    draw_box(ax, (0.04, 0.48), (0.17, 0.25), "候选状态 s_t", "38维\n全局 / 类别 / 轮次", "#2D70AA")
    draw_box(ax, (0.27, 0.48), (0.17, 0.25), "State Encoder", "38 → 128 → 64\nLayerNorm + ReLU", "#2D70AA")
    draw_box(ax, (0.04, 0.14), (0.17, 0.25), "提示动作 a", "20维\n类别 / rank / role / overlap", "#17857D")
    draw_box(ax, (0.27, 0.14), (0.17, 0.25), "Action Encoder", "20 → 64 → 64\nLayerNorm + ReLU", "#17857D")
    draw_box(ax, (0.52, 0.31), (0.19, 0.28), "交互融合", "[e_s ; e_a ; e_s*e_a]\n192维", "#D99124")
    draw_box(ax, (0.78, 0.31), (0.17, 0.28), "Action Scorer", "192 → 96 → 1\nargmax / STOP", "#C65049")
    arrow(ax, (0.21, 0.605), (0.26, 0.605), "#2D70AA")
    arrow(ax, (0.21, 0.265), (0.26, 0.265), "#17857D")
    arrow(ax, (0.44, 0.605), (0.51, 0.50), "#718096")
    arrow(ax, (0.44, 0.265), (0.51, 0.39), "#718096")
    arrow(ax, (0.71, 0.45), (0.77, 0.45), "#D99124")
    fig.tight_layout(pad=0.4)
    path2 = ASSET_DIR / "dynamic_policy_network.png"
    fig.savefig(path2, bbox_inches="tight", facecolor="#F6F7F9")
    plt.close(fig)
    return path, path2


def add_method_chapter(doc: Document, anchor, overall_fig: Path, policy_fig: Path):
    add_p(anchor, "引言", SECTION_STYLE)
    add_p(anchor, "航拍图像中的目标尺度小、分布密集且类别外观相近，单次开放词表检测容易同时出现漏检与高分误检。固定提示词虽然便于执行，却难以适应自然语言查询中的属性、关系和同义表达；无差别执行全部扩展提示又会造成候选框数量和计算开销快速增长。针对这一矛盾，本文将最新系统组织为“语义扩展—策略选择—候选融合—排序校准”的闭环流程。")
    add_p(anchor, "具体而言，Qwen2.5-VL-3B-Instruct 根据输入 query 或 caption 生成当前图像专属的短检测提示词集合；强化学习策略利用当前候选池的统计状态，对 STOP 与各个动态提示动作逐一评分；被选择的提示词直接驱动 Grounding DINO 生成补充候选，经合并和同类去重后形成下一轮状态；循环结束后，Exp38 排序器利用候选来源、查询匹配、几何属性和多路分数等富元信息重标定最终置信度。该设计将语言模型的开放语义能力、强化学习的预算分配能力和监督排序器的置信度校准能力分别置于最合适的环节。")
    add_p(anchor, "本章首先介绍总体框架，然后说明 Qwen 动态提示词池、强化学习动态动作选择、候选合并与 Exp38 排序器，最后给出训练和推理流程。")

    add_p(anchor, "方法总体框架", SECTION_STYLE)
    add_p(anchor, "给定航拍图像 I 与文本查询 q，系统输出类别、坐标和校准分数组成的检测集合。整体映射记为：")
    add_formula(anchor, "Y = F_rank(F_loop(F_det(I, q, P(q))))，                         （3-1）")
    add_p(anchor, "其中，P(q) 表示 Qwen 生成的动态提示词池，F_det 表示 Grounding DINO 候选生成，F_loop 表示强化学习控制的多轮候选扩展，F_rank 表示 Exp38 富元信息排序。与将强化学习用于逐框保留或删除的旧方案不同，最新版本中的 RL 决策对象是“下一轮是否执行某条提示词”，其作用是控制候选生成动作和计算预算；候选框的最终优先级由 Exp38 排序器负责。")
    add_picture(doc, anchor, overall_fig, "图3-1 Qwen 动态提示词、RL 闭环选择与 Exp38 排序总体流程", 6.4)
    add_p(anchor, "推理从 query-aware 初始候选集合 C0 开始。在第 t 次决策时，系统从 Ct 提取状态 st，并为尚未使用的动态提示词构造动作特征。若策略选择 STOP，则不再增加候选；若选择提示动作 at，则用该提示调用 Grounding DINO，得到补充候选 ΔCt，并执行合并与去重。三轮设置包含第一轮初始候选和后续两次反馈决策。循环结束后，所有候选统一进入 Exp38 重评分和类别感知后处理。")

    add_p(anchor, "Qwen 动态提示词与动作空间生成", SECTION_STYLE)
    add_p(anchor, "查询解析与结构化提示词池", SUBSECTION_STYLE)
    add_p(anchor, "本文使用本地 Qwen2.5-VL-3B-Instruct 作为提示词提议器，而不让语言模型直接输出检测框或最终类别。模型接收自然语言查询，输出最多 8 个结构化对象，每个对象包含 VisDrone 目标类别和一条简短英文检测短语。类别被约束在 pedestrian、people、bicycle、car、van、truck、tricycle、awning tricycle、bus 和 motor 十类中，从而保证开放表达能够落到统一评价空间。")
    add_formula(anchor, "P(q) = Qwen(q) = {(c_k, p_k)}_(k=1)^K，K ≤ 8。             （3-2）")
    add_p(anchor, "提示词不仅保留类别名称，还可包含颜色、服饰、位置关系和场景动作。例如查询“A black-clad cyclist is crossing the intersection”可被改写为“a cyclist in dark clothing moving through an intersection”等短语。生成结果经规范化、空项过滤和重复项删除后构成当前图像的 prompt pool；若生成失败，则使用覆盖主要弱类的回退提示集合，以保证推理可继续执行。")
    add_p(anchor, "提示动作的结构化表示", SUBSECTION_STYLE)
    add_p(anchor, "每个提示动作包含文本 text、目标类别 class、语言模型输出次序 llm_rank 与语义角色 llm_role。语义角色分为 primary、alias、attribute、relation 和 fallback，用于区分核心类别表达、同义表达、属性描述、关系描述和回退提示。动态动作集合定义为：")
    add_formula(anchor, "A_t = {STOP} ∪ {a_k | p_k ∈ P(q), p_k 尚未执行}。          （3-3）")
    add_p(anchor, "与固定 12 动作策略不同，动作数量和内容随查询变化。测试缓存中 300 个样本共生成 2396 条提示词，其中唯一提示词 604 条，说明动作空间并非简单复用固定类别名称。最新实现会将被选中的原始提示词直接送入 Grounding DINO 生成候选，而不是仅映射到已有离线候选来源。")
    add_table(doc, anchor, "表3-1 动态提示动作的 20 维特征组成",
              ["特征组", "维数", "内容"],
              [["STOP 标记", "1", "区分终止动作与扩展动作"],
               ["类别 one-hot", "10", "对应 VisDrone 十类"],
               ["提示统计", "5", "词数、LLM rank、query 词重合率、轮次、是否已使用"],
               ["语义角色", "4", "primary、alias、attribute、relation"],
               ["合计", "20", "作为 Action Encoder 输入"]],
              [1.55, 0.65, 4.0], (4,))

    add_p(anchor, "强化学习驱动的闭环候选扩展", SECTION_STYLE)
    add_p(anchor, "38维候选状态", SUBSECTION_STYLE)
    add_p(anchor, "状态向量描述当前候选池而不是单个候选框。其前 6 维为全局统计，包括归一化候选数量、平均分、最高分、分数标准差、平均面积比例和弱小类别候选比例；随后按十个类别分别统计候选数量、最高分和平均面积，共 30 维；最后 2 维编码当前进度和剩余动作预算，因此状态总维数为 38。")
    add_formula(anchor, "s_t = [g_t^(6), n_t^(10), m_t^(10), a_t^(10), progress_t, remaining_t] ∈ R^38。    （3-4）")
    add_p(anchor, "这种状态构造能够同时反映候选池规模、置信度分布、尺度构成和类别缺口。例如某一弱类候选数量少且最高分低时，相关扩展提示可能获得更高价值；当候选已经充分或预算不足时，STOP 动作应更具优势。")
    add_p(anchor, "状态—动作匹配策略网络", SUBSECTION_STYLE)
    add_p(anchor, "固定动作策略直接由状态输出固定维数 logits，无法为未见过的新提示词打分。为适配动态 prompt pool，本文采用双编码器结构。State Encoder 将 38 维状态映射为 64 维状态嵌入，Action Encoder 将 20 维动作特征映射为 64 维动作嵌入；随后拼接两者及其逐元素乘积，经 MLP 输出该状态下执行该动作的标量得分：")
    add_formula(anchor, "Q(s_t,a_k) = MLP([E_s(s_t); E_a(a_k); E_s(s_t) ⊙ E_a(a_k)])。          （3-5）")
    add_p(anchor, "State Encoder 的结构为 38→128→64，Action Encoder 为 20→64→64，交互向量维数为 192，Action Scorer 为 192→96→1。推理时对当前有效动作逐一计算分数并选择最大者；已经使用的提示词通过动作掩码移除，避免重复调用。")
    add_picture(doc, anchor, policy_fig, "图3-2 动态状态—动作匹配策略网络", 6.25)
    add_p(anchor, "奖励函数与策略训练", SUBSECTION_STYLE)
    add_p(anchor, "训练样本由验证集候选状态、动态提示动作和执行后的检测收益组成。奖励同时考虑普通召回、高 IoU 召回、小目标召回和精度增量，并对新增候选数量、动作次数和不恰当的提前停止施加代价：")
    add_formula(anchor, "R = 0.40ΔRecall + 0.22ΔRecall@0.75 + 0.45ΔSmallRecall + 0.15ΔPrecision")
    add_formula(anchor, "    − 0.010 Cost_box − 0.012 N_action − 0.012 I_stop。                    （3-6）")
    add_p(anchor, "对于同一状态下的候选动作组，先将真实奖励按温度系数转化为软目标分布，再用交叉熵项、期望奖励项和熵正则联合优化策略。该训练方式并非让 Qwen 参与参数更新，而是学习在给定候选状态下如何评价 Qwen 提出的动作。最佳动态策略在约第 50 个 epoch 获得验证集平均奖励 0.028201，状态维数和动作特征维数分别为 38 和 20。")
    add_p(anchor, "多轮闭环推理", SUBSECTION_STYLE)
    add_p(anchor, "每轮若选择扩展动作，系统直接使用对应提示词调用 Grounding DINO，并将新候选与当前候选池合并。对于同类别且 IoU 不低于 0.92 的高度重复框，仅保留分数更高者：")
    add_formula(anchor, "C_(t+1) = Dedup(C_t ∪ F_GDINO(I,p_(a_t)))。                 （3-7）")
    add_p(anchor, "若选择 STOP，则 Ct+1= Ct。默认三轮推理包含 C0 和两次策略决策。闭环轮数并非越多越好：更多轮次可补充有效候选，但也会增加动作调用和候选规模，因此第 4 章将对 2、3、4 轮进行同策略消融。")

    add_p(anchor, "Exp38 富元信息排序器", SECTION_STYLE)
    add_p(anchor, "候选匹配与富元信息构建", SUBSECTION_STYLE)
    add_p(anchor, "RL 解决“是否继续生成以及使用什么提示生成”的问题，Exp38 进一步解决“哪些候选应排在前面”的问题。对于闭环输出候选，系统按类别和 IoU 与在线 metadata 记录匹配，并构造固定 68 维富元信息。特征包括多路置信度及交互项、面积与宽高等几何量、query 匹配分数、source_prompt、source_type、source_class_id、类别与尺寸先验、候选 rank 和类别 one-hot 等。")
    add_table(doc, anchor, "表3-2 Exp38 富元信息的主要组成",
              ["特征组", "代表字段", "作用"],
              [["分数及交互", "final/gdino/query/alias 等分数、差值与乘积", "描述多路模型响应及一致性"],
               ["几何与排序", "面积、宽高、宽高比、中心、候选 rank", "描述尺度、位置和初始排序"],
               ["候选来源", "source_prompt、source_type、source_class_id", "建模不同提示来源的系统偏差"],
               ["类别与先验", "类别 one-hot、尺寸层级、小目标先验", "提供航拍类别与尺度信息"],
               ["完整输入", "68维固定槽位", "输入 RichFusionModel"]],
              [1.25, 2.75, 2.35], (4,))
    add_p(anchor, "排序模型与分数融合", SUBSECTION_STYLE)
    add_p(anchor, "Exp38 由别名候选重排分支与 RichFusionModel 组成。别名分支重点区分弱小类别真阳性与 hard negative；RichFusionModel 对 68 维标准化特征进行 68→128→80→1 的非线性映射，并使用 LayerNorm、ReLU 和 Dropout。最终分数由基础融合分数与富元信息预测分数线性组合，融合系数只在验证集 holdout 上选择，测试集不参与调参。")
    add_formula(anchor, "score_i^final = (1−α) score_i^base + α score_i^fusion。              （3-8）")
    add_p(anchor, "排序训练使用 class-aware IoU 匹配产生候选框级监督标签，同类且满足阈值的候选记为正样本，其余候选作为负样本或 hard negative。模型只重标定分数，不改变候选坐标和类别，因此固定候选集上的 mAP 变化可直接用于衡量排序质量。")
    add_p(anchor, "类别感知后处理", SUBSECTION_STYLE)
    add_p(anchor, "重评分后，系统按照类别分别执行 NMS 和 Top-K 约束，避免不同类别之间相互抑制。该步骤删除高度重叠的同类框并控制每类输出规模，但不承担语义扩展或策略决策功能。至此，Qwen、RL 与排序器形成清晰分工：Qwen 提议可执行语义动作，RL 分配扩展预算，Grounding DINO 生成候选，Exp38 校准候选排序。")

    add_p(anchor, "训练与推理流程", SECTION_STYLE)
    add_p(anchor, "训练阶段分为两步。第一步在验证集上为 Qwen 生成的各提示动作实际执行候选扩展，根据 GT 匹配统计构造奖励表，训练动态状态—动作评分网络；第二步在固定训练/holdout 划分上训练 Exp38，并仅用 holdout 选择融合系数。两类模型均冻结 Grounding DINO 和 Qwen 主体参数，避免高成本端到端微调。")
    add_p(anchor, "推理阶段不使用 GT 和奖励计算。系统先由 Qwen 生成提示词池，再从 query-aware C0 开始执行最多两次动态策略决策；每次扩展后更新 38 维状态；最后对合并候选执行 Exp38 重评分、类别感知 NMS 和 Top-K。Qwen 提示词可预先缓存以提高批量评测的可复现性，但实际选中提示所对应的 Grounding DINO 候选由提示词直接生成。")
    add_p(anchor, "当前方法仍有两点边界。第一，Qwen 的类别字段被约束在 VisDrone 十类，开放文本表达更丰富，但最终评价类别空间仍是封闭的；第二，动态策略在候选效率与召回之间形成不同偏好，Qwen 动态策略更保守，而固定动作策略具有更高召回。第 4 章将据此分别报告排序质量、召回和动作开销，避免用单一指标概括所有方法。")

    add_p(anchor, "本章小结", SECTION_STYLE)
    add_p(anchor, "本章介绍了最新的 Qwen 动态提示词、强化学习候选扩展和 Exp38 排序器联合方法。Qwen2.5-VL-3B-Instruct 将自然语言查询转换为图像相关的结构化提示词池；动态策略网络通过 38 维候选状态和 20 维动作特征学习状态—动作匹配，在 STOP 与提示扩展动作之间进行多轮选择；被选提示直接驱动 Grounding DINO 生成补充候选；Exp38 再利用 68 维富元信息校准候选置信度。该方法将候选覆盖、动作预算和最终排序三个问题解耦，为后续实验中的轮数、动作策略和元数据特征消融提供了明确的控制变量。")


def add_experiment_chapter(doc: Document, anchor):
    add_p(anchor, "引言", SECTION_STYLE)
    add_p(anchor, "本章围绕最新的“Qwen 动态提示词 + RL + Exp38 排序器”流程开展实验。实验重点回答四个问题：动态提示词策略与固定动作策略相比具有何种精度—开销权衡；增加闭环轮数是否持续有效；RL 的收益是否仅来自执行更多候选生成动作；Exp38 的提升是否确实来自富元信息而非候选集合或模型容量变化。")
    add_p(anchor, "为避免混用不同候选来源和旧 checkpoint，本章优先采用答辩 PPT 中的统一测试集结果及 2026 年 9 月补充消融。早期不同数据划分、不同动作空间或缺少完整复现产物的记录不进入正式横向比较。")

    add_p(anchor, "实验设置", SECTION_STYLE)
    add_p(anchor, "数据集与评价协议", SUBSECTION_STYLE)
    add_p(anchor, "主实验使用 VisDroneSplit1000Guarded 的固定 test 划分，共 150 张图像和 7973 个 GT 框，类别包括 pedestrian、people、bicycle、car、van、truck、tricycle、awning tricycle、bus 和 motor。RefDrone 与 AerialVG 的真实语言短语用于补充查询表达，不与主测试集指标混合。所有对比采用 class-aware 一对一匹配：预测类别与 GT 类别一致且 IoU 达到阈值时才计为命中。")
    add_table(doc, anchor, "表4-1 数据集与统一实验协议",
              ["项目", "设置", "说明"],
              [["主测试集", "VisDroneSplit1000Guarded/test", "150张图像，7973个GT框"],
               ["类别空间", "VisDrone 10类", "Qwen 输出类别受统一类别表约束"],
               ["候选生成", "Grounding DINO", "选中动态提示后直接生成候选"],
               ["匹配方式", "class-aware，IoU≥0.5/0.75", "同类一对一匹配"],
               ["公平性控制", "相同 C0、去重、Exp38 与评测脚本", "动作策略消融保持其余流程一致"]],
              [1.25, 2.35, 2.75])
    add_p(anchor, "评价指标与实现细节", SUBSECTION_STYLE)
    add_p(anchor, "总体指标采用 Acc@0.5、Acc@0.75 和 mAP@0.5，并报告预测框数量、TP@0.5、每图平均非 STOP 动作数和 STOP 比例。小目标按像素面积不超过 32×32 定义；Small Candidate Recall@0.5 表示 Exp38 重评分前候选对小目标 GT 的覆盖率，Small-object Final Recall@0.5 表示重评分后最终结果对同一小目标集合的召回率。")
    add_p(anchor, "动态策略使用 Qwen2.5-VL-3B-Instruct，策略状态维数 38、动作特征维数 20。固定轮数实验的时间统计仅包含缓存候选读取、策略决策、候选合并去重和 Exp38 重评分，不包含 Grounding DINO 原始候选生成及指标评测，因此该时间用于同口径轮数比较，不代表完整在线端到端时延。")

    add_p(anchor, "Qwen 动态动作与固定策略对比", SECTION_STYLE)
    add_p(anchor, "表4-2比较只使用第一轮 query-aware 候选的 Query base、固定动作三轮 RL 和 Qwen 动态动作三轮 RL。固定动作策略从预定义弱类扩展动作中选择；动态策略针对每张图的 Qwen prompt pool 评分。")
    add_table(doc, anchor, "表4-2 Qwen 动态动作与固定策略对比",
              ["方法", "预测框数", "Acc@0.5", "Acc@0.75", "mAP@0.5"],
              [["Query base + Exp38", "22,298", "0.4651", "0.3131", "0.1343"],
               ["固定动作三轮 RL", "40,904", "0.5293", "0.3310", "0.1397"],
               ["Qwen 动态动作三轮 RL", "35,869", "0.5004", "0.3248", "0.1424"]],
              [2.05, 1.05, 1.0, 1.0, 1.0], (1, 2))
    add_p(anchor, "固定动作 RL 的 Acc@0.5 最高，为 0.5293，说明针对弱类的固定扩展具有更强召回能力。Qwen 动态策略将预测框从 40904 减少到 35869，降幅约 12.3%，同时将 mAP@0.5 从 0.1397 提升到 0.1424；但 Acc@0.5 和 Acc@0.75 分别下降 0.0289 和 0.0062。动作统计中，动态策略的 STOP 比例为 52.3%，高于固定策略的 32.3%，表明其选择更保守。")
    add_p(anchor, "因此，本实验不支持“动态策略在所有指标上优于固定策略”的结论。更准确的表述是：Qwen 动态动作改善了候选规模和置信度排序质量，固定动作策略则更有利于候选召回；两者体现了不同的精度—开销偏好。")

    add_p(anchor, "闭环轮数消融", SECTION_STYLE)
    add_p(anchor, "轮数消融固定同一三轮训练策略，仅改变总推理轮数。两轮表示 C0 后执行一次反馈决策，三轮执行两次；四轮沿用同一策略，并在额外决策处使用 progress=1.0、remaining=0.0 的范围终点编码，因此它是固定策略下的推理轮数扩展，而非重新训练的四轮策略。")
    add_table(doc, anchor, "表4-3 RL 闭环轮数消融结果",
              ["总轮数", "候选框", "TP@0.5", "Acc@0.5", "Acc@0.75", "mAP@0.5", "单图时间"],
              [["2", "37,487", "4,097", "0.5139", "0.3253", "0.1386", "75.11 ms"],
               ["3", "40,904", "4,220", "0.5293", "0.3310", "0.1397", "93.16 ms"],
               ["4", "42,229", "4,271", "0.5357", "0.3335", "0.1402", "106.11 ms"]],
              [0.62, 0.9, 0.85, 0.9, 0.9, 0.9, 1.0], (1, 2))
    add_p(anchor, "四轮取得最高最终指标，但边际收益递减。由三轮增加到四轮仅新增 51 个 TP@0.5，Acc@0.5 提升 0.0064，mAP@0.5 提升 0.00052，同时增加 1325 个候选框和 1.94 s 闭环处理时间。因此最高召回优先时可使用四轮，综合候选规模和效率时三轮更均衡。")

    add_p(anchor, "RL 动作策略消融", SECTION_STYLE)
    add_p(anchor, "为判断收益是否只是因为多执行了候选生成，本实验在相同第一轮候选、去重、Exp38 重评分和评测逻辑下比较立即停止、穷举扩展、随机动作、Qwen 顺序选择、固定动作 RL 与 Qwen 动态 RL。Random action 为随机种子 42—46 的均值与总体标准差；其每图非 STOP 动作预算与固定动作 RL 一致。")
    add_table(doc, anchor, "表4-4 RL 动作策略消融结果",
              ["方法", "预测框数", "Acc@0.5", "mAP@0.5", "动作/图", "STOP", "小目标最终Recall"],
              [["No expansion", "22,298", "0.4651", "0.1343", "0.000", "100.0%", "0.3291"],
               ["All expansion", "68,771", "0.5731", "0.1510", "11.000", "0.0%", "0.4593"],
               ["Random (5 seeds)", "30,227±582", "0.4892±0.0034", "0.1377±0.0012", "1.353", "0.0%", "0.3584±0.0053"],
               ["Prompt-rank", "41,969", "0.5090", "0.1440", "1.087", "0.0%", "0.3840"],
               ["Fixed-action RL", "40,904", "0.5293", "0.1397", "1.353", "32.3%", "0.4159"],
               ["Qwen dynamic RL", "35,869", "0.5004", "0.1424", "0.953", "52.3%", "0.3732"]],
              [1.45, 0.9, 1.08, 1.03, 0.75, 0.72, 1.15], (4, 5))
    add_p(anchor, "在相同的 1.353 动作/图预算下，固定动作 RL 的 Acc@0.5 比随机动作高 4.01 个百分点，小目标最终 Recall@0.5 高 5.75 个百分点，说明策略确实学到了更有效的动作选择，而不仅是增加生成次数。All expansion 的绝对指标最高，但每图 11 个动作，是固定动作 RL 的 8.13 倍，最终框数也是其 1.68 倍。由此可见，RL 的主要价值是提高给定动作和候选预算下的扩展效率，而不是证明穷举无法提高精度。")
    add_p(anchor, "Prompt-rank 的 mAP@0.5 为 0.1440，高于固定动作 RL；Qwen dynamic RL 用更少动作获得 0.1424。动态方法在排序质量和候选规模上具有优势，而固定 RL 的召回和小目标召回更强。")
    if ACTION_FIG.exists():
        add_picture(doc, anchor, ACTION_FIG, "图4-1 不同动作策略的主要指标与候选框数量", 6.3)

    add_p(anchor, "Exp38 固定候选集排序实验", SECTION_STYLE)
    add_p(anchor, "为验证排序收益不依赖候选增加，三种方法读取同一来源的 52872 个 metadata 候选框。规则筛选对照只删除候选，排序器校准只重标定分数。表4-5为最终单模型配置结果。")
    add_table(doc, anchor, "表4-5 固定候选集上的排序优化结果",
              ["方法", "预测框数", "Acc@0.5", "Acc@0.75", "mAP@0.5", "TP@0.5"],
              [["富元信息融合前", "52,872", "0.5374", "0.3463", "0.1464", "4,285"],
               ["规则筛选对照", "9,586", "0.3193", "0.2491", "0.1187", "2,546"],
               ["排序器校准", "52,872", "0.5374", "0.3491", "0.1690", "4,285"]],
              [1.65, 1.0, 0.95, 0.95, 0.95, 0.9], (2,))
    add_p(anchor, "排序器校准保持预测框数、Acc@0.5 和 TP@0.5 不变，将 Acc@0.75 从 0.3463 提升到 0.3491，mAP@0.5 从 0.1464 提升到 0.1690。由于坐标、类别和候选集合均未改变，收益来自置信度校准与排序。规则硬筛选将候选压缩到 9586 个，但同时显著损失召回和 mAP，说明仅依赖固定规则删除候选难以替代学习式排序。")

    add_p(anchor, "Exp38 元数据特征消融", SECTION_STYLE)
    add_p(anchor, "进一步的严格消融固定 52872 个候选、68 维输入槽位和 19649 个模型参数，仅将未启用的特征位置置零。每组使用随机种子 42、43、44 训练 60 epochs，融合系数只在 val holdout 上选择，test 仅用于最终评价。六组配置采用逐步累积设计。")
    add_table(doc, anchor, "表4-6 Exp38 元数据特征累积消融（均值±样本标准差）",
              ["配置", "启用维数", "Acc@0.5", "Acc@0.75", "mAP@0.5", "逐步ΔmAP"],
              [["Raw score", "1", "0.5374±0.0000", "0.3463±0.0000", "0.1464±0.0000", "+0.0000"],
               ["+ geometry", "9", "0.5374±0.0000", "0.3463±0.0005", "0.1504±0.0012", "+0.0040"],
               ["+ query", "10", "0.5374±0.0000", "0.3461±0.0002", "0.1514±0.0005", "+0.0010"],
               ["+ source", "36", "0.5374±0.0000", "0.3483±0.0002", "0.1664±0.0003", "+0.0150"],
               ["+ small prior", "59", "0.5374±0.0000", "0.3482±0.0003", "0.1637±0.0008", "−0.0027"],
               ["Full Exp38", "68", "0.5374±0.0000", "0.3490±0.0006", "0.1677±0.0020", "+0.0040"]],
              [1.2, 0.72, 1.12, 1.12, 1.12, 0.95], (3, 5))
    add_p(anchor, "Full Exp38 相比 score-only MLP 的 mAP@0.5 绝对提升 0.0212，相对提升 14.5%。所有配置的候选框数量和 Acc@0.5 不变，且网络结构和参数量相同，因此差异来自元数据字段对排序的贡献。加入 source 特征后 mAP 从 0.1514 提升至 0.1664，是最大且最稳定的单步增益；几何和 query 特征提供较小正增益。")
    add_p(anchor, "加入 small prior 后 mAP 从 0.1664 降到 0.1637，不能将该先验单独表述为有效模块，可能原因是过拟合或与来源特征重复。完整分数及交互特征将 mAP 恢复至最高均值 0.1677。需要说明，表4-5的 0.1690 是最终单模型配置结果，而表4-6的 0.1677±0.0020 是三随机种子、等参数严格控制消融的均值，两者实验目的和统计口径不同，不应直接合并为同一个数值。")

    add_p(anchor, "定性分析与局限性", SECTION_STYLE)
    add_p(anchor, "在典型样例中，输入查询描述穿过路口的黑衣骑行者。Qwen 生成 7 条包含 cyclist、dark clothing、intersection 和 road 等信息的短提示；策略从 STOP 与动态动作中依次选择 person 和 bicycle。候选数量由 C0 的 291 个增加到 C1 的 642 个，最终合并后为 692 个，其中 483 个候选在 Exp38 metadata 中完成高 IoU 匹配和重评分。该流程说明语言扩展并不直接决定结果，而是通过 RL 的预算选择与 Exp38 的排序校准共同作用。")
    add_p(anchor, "当前误差主要来自三方面。第一，动态策略较高的 STOP 比例会减少候选，但也可能提前终止弱目标扩展；第二，开放提示仍受 VisDrone 十类输出空间约束，car、van、truck 和 bus 等俯视外观相近类别仍易混淆；第三，当 Grounding DINO 对极小目标完全没有响应时，后续排序只能重排已有候选，不能凭空恢复目标。后续可联合学习停止代价和类别特定动作预算，并为车辆细粒度类别引入更强视觉特征。")

    add_p(anchor, "本章小结", SECTION_STYLE)
    add_p(anchor, "本章按照统一测试集和控制变量验证了最新流程。Qwen 动态动作三轮策略以 35869 个预测框取得 0.1424 mAP@0.5，相比固定动作策略减少约 12.3% 候选并改善 mAP，但固定策略在 Acc 和小目标召回上更高。轮数实验表明四轮指标最高、三轮效率更均衡；动作消融证明固定 RL 在相同动作预算下显著优于随机选择，而穷举扩展虽有最高绝对指标，却需要 8.13 倍动作预算。")
    add_p(anchor, "固定候选集实验进一步表明，Exp38 的收益来自排序而不是候选增加。最终单模型将 mAP@0.5 从 0.1464 提升到 0.1690；三随机种子严格元数据消融中，Full Exp38 达到 0.1677±0.0020，较 Raw score 相对提升 14.5%，其中 source 信息贡献最大。综上，最新系统的三类模块形成互补关系：Qwen 提供开放语义动作，RL 提高候选扩展的预算效率，Exp38 改善最终置信度排序。")


def set_update_fields(doc: Document):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def replace_internal_experiment_name(doc: Document):
    replacements = [
        ("Full Exp38", "完整特征配置"),
        ("Exp38 排序总体流程", "富元信息融合排序总体流程"),
        ("Exp38 富元信息的主要组成", "富元信息融合排序器的特征组成"),
        ("Exp38 固定候选集排序实验", "固定候选集排序实验"),
        ("Exp38 元数据特征累积消融", "富元信息特征累积消融"),
        ("Exp38 元数据特征消融", "富元信息特征消融"),
        ("Query base + Exp38", "Query base + 富元信息排序"),
        ("Exp38 由", "富元信息融合排序器由"),
        ("Exp38 再", "富元信息融合排序器再"),
        ("训练 Exp38", "训练富元信息融合排序器"),
        ("Exp38 的", "富元信息融合排序器的"),
        ("Exp38 改善", "富元信息融合排序器改善"),
        ("Exp38 富元信息排序器", "富元信息融合排序器"),
        ("Exp38 富元信息排序", "富元信息融合排序"),
        ("Exp38 排序器", "富元信息融合排序器"),
        ("Exp38 重评分", "富元信息重评分"),
        ("Exp38 metadata", "排序器 metadata"),
        ("Exp38", "富元信息融合排序器"),
        ("EXP38", "富元信息融合排序器"),
        ("实验38", "富元信息融合排序器"),
    ]

    def replace_runs(paragraph):
        for run in paragraph.runs:
            for old, new in replacements:
                if old in run.text:
                    run.text = run.text.replace(old, new)

    for paragraph in doc.paragraphs:
        replace_runs(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_runs(paragraph)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                replace_runs(paragraph)


def main():
    overall_fig, policy_fig = make_framework_figures()
    doc = Document(SOURCE)
    chapter3_index = next(i for i, p in enumerate(doc.paragraphs)
                          if p.style.name == "Heading 1" and p.text.strip() == "强化学习驱动的闭环航拍图像识别方法")
    original_before_method = [p.text for p in doc.paragraphs[:chapter3_index]]

    chapter3 = find_heading(doc, "强化学习驱动的闭环航拍图像识别方法")
    chapter4 = find_heading(doc, "实验设置与结果分析")
    conclusion = find_paragraph(doc, "结论")
    remove_between(chapter3, chapter4)
    remove_between(chapter4, conclusion)
    chapter3.text = "Qwen动态提示词与强化学习驱动的闭环航拍图像识别方法"
    add_method_chapter(doc, chapter4, overall_fig, policy_fig)
    add_experiment_chapter(doc, conclusion)
    replace_internal_experiment_name(doc)
    set_update_fields(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    check = Document(OUT)
    new_chapter3_index = next(i for i, p in enumerate(check.paragraphs)
                              if p.style.name == "Heading 1" and p.text.strip() == "Qwen动态提示词与强化学习驱动的闭环航拍图像识别方法")
    current_before_method = [p.text for p in check.paragraphs[:new_chapter3_index]]
    if original_before_method != current_before_method:
        raise RuntimeError("Content before Chapter 3 changed unexpectedly")
    print(OUT)
    print(f"paragraphs={len(check.paragraphs)} tables={len(check.tables)} sections={len(check.sections)}")


if __name__ == "__main__":
    main()
